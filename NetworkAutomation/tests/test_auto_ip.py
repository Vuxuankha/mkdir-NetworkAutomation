"""Offline tests; all network operations are simulated. Never uses production DB."""
import importlib
import json
import sys
import tempfile
import threading
import time
import unittest
from dataclasses import replace
from pathlib import Path
from unittest.mock import patch

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from database import db
from modules import auto_ip as a
from openpyxl import Workbook, load_workbook

USER = {'username': 'offline-test', 'role': 'Admin'}


class FakeAdapter:
    active = 0
    max_active = 0
    seen = []
    lock = threading.Lock()

    def __init__(self, host, credential, options, stop):
        self.host, self.credential, self.options, self.stop = host, credential, options, stop
        self.secret_values = [credential.get('community', '')]
        with self.lock:
            type(self).active += 1
            type(self).max_active = max(type(self).max_active, type(self).active)
            type(self).seen.append(host)

    def check(self):
        if self.stop.is_set():
            raise a.Cancelled('Stopped')

    def ping(self):
        self.stop.wait(0.01)
        self.check()
        return {'alive': False, 'latency_ms': None, 'packet_loss': 100}

    def tcp(self):
        self.check()
        return {'22': 'OPEN', '443': 'OPEN'}

    def get(self, oids):
        self.check()
        if self.credential['mode'] == 'off':
            raise a.SkipTask('No SNMP profile')
        if self.host.endswith('.2'):
            raise TimeoutError('Authentication failure for ' + self.credential.get('community', ''))
        result = {}
        for oid in oids:
            values = {a.SYS['name']: 'lab-' + self.host, a.SYS['descr']: 'Cisco IOS XE Software',
                      a.SYS['object']: '1.3.6.1.4.1.9.1.1234', a.SYS['uptime']: 1000000,
                      '1.3.6.1.4.1.9.2.1.58.0': 35}
            result[oid] = values.get(oid)
            for key, base in a.IF_OIDS.items():
                if oid == base + '.1':
                    result[oid] = {'name': 'Ethernet1', 'oper': 1, 'speed': 1000000000,
                                   'in32': 1000, 'out32': 2000, 'in64': 1000000, 'out64': 2000000,
                                   'in_errors': 0, 'out_errors': 0, 'in_discards': 0,
                                   'out_discards': 0, 'highspeed': 1000}[key]
        return result

    def walk(self, base, limit=None):
        self.check()
        if base == a.IF_INDEX:
            return {base + '.1': 1}, False
        if base == '1.0.8802.1.1.2.1.4.1.1.9':
            return {base + '.0.1.1': 'neighbor'}, False
        if base == '1.0.8802.1.1.2.1.4.1.1.7':
            return {base + '.0.1.1': 'Ethernet2'}, False
        return {}, False

    def backup(self, driver):
        raise a.SkipTask('No SSH credential')

    def close(self):
        with self.lock:
            type(self).active -= 1


class AutoIPTests(unittest.TestCase):
    def setUp(self):
        self.temp = tempfile.TemporaryDirectory(prefix='nms-autoip-test-')
        self.path = Path(self.temp.name)
        self.patches = []
        self.engine = None
        modules = [db, a]
        for name in ['advanced_pages', 'nms_v3', 'nms_v4', 'nms_v5', 'nms_v6', 'nms_v7', 'nms_v8', 'nms_v9', 'nms_v10', 'nms_v11', 'nms_v12']:
            modules.append(importlib.import_module('modules.' + name))
        for mod in modules:
            for key, value in [('DB_PATH', self.path / 'test.db'), ('APP_DIR', self.path),
                               ('KEY_FILE', self.path / 'key'), ('BACKUP_DIR', self.path / 'backups')]:
                if hasattr(mod, key):
                    p = patch.object(mod, key, value)
                    p.start()
                    self.patches.append(p)
        db.init_database()
        import modules.advanced_pages as adv
        adv.ensure_advanced_tables()
        for v in range(3, 13):
            mod = importlib.import_module(f'modules.nms_v{v}')
            getattr(mod, f'ensure_v{v}_tables')()
        a.ensure_tables()
        a.save_profile('Lab', 'v2c', 'never-print-secret', role='Admin')
        FakeAdapter.seen, FakeAdapter.active, FakeAdapter.max_active = [], 0, 0

    def tearDown(self):
        if self.engine:
            self.engine.stop()
            if self.engine.thread:
                self.engine.thread.join(10)
        for p in reversed(self.patches):
            p.stop()
        self.temp.cleanup()

    def run_job(self, targets=None, **kwargs):
        options = replace(a.Options(default_profile='Lab', report=False, alerts=False), **kwargs)
        self.engine = a.AutomationEngine(FakeAdapter)
        self.engine.start(targets or [a.Target('192.0.2.1')], options, USER, authorized=True)
        self.engine.thread.join(12)
        self.assertFalse(self.engine.running)
        self.assertEqual(self.engine.snapshot()['status'], 'Completed', self.engine.snapshot())
        return a.run_steps(self.engine.snapshot()['run_id'], limit=10000)

    def test_ip_split_deduplicate(self):
        targets, duplicates = a.parse_text('192.0.2.1, 192.0.2.2;\n192.0.2.1')
        self.assertEqual(len(targets), 2)
        self.assertEqual(duplicates, 1)

    def test_ipv6_canonical(self):
        targets, _ = a.parse_text('2001:0db8:0:0:0:0:0:1')
        self.assertEqual(targets[0].ip, '2001:db8::1')

    def test_invalid_ip_and_ranges(self):
        for value in ['192.0.2.0/24', 'example.com', '0.0.0.0', '224.1.1.1', '255.255.255.255', 'fe80::1%eth0']:
            with self.assertRaises(ValueError):
                a.parse_text(value)

    def test_empty_rejected(self):
        with self.assertRaises(ValueError):
            a.parse_text('')

    def workbook(self, rows, name='input.xlsx'):
        wb = Workbook()
        for row in rows:
            wb.active.append(row)
        path = self.path / name
        wb.save(path)
        wb.close()
        return path

    def test_excel_minimal(self):
        result, _ = a.parse_excel(self.workbook([['IP'], ['192.0.2.1']]))
        self.assertEqual(result[0].ip, '192.0.2.1')

    def test_excel_optional_vietnamese_headers(self):
        rows = [['\u0110\u1ecba ch\u1ec9 IP', 'T\u00ean thi\u1ebft b\u1ecb', 'H\u1ed3 s\u01a1'], ['192.0.2.1', 'Switch', 'Lab']]
        result, _ = a.parse_excel(self.workbook(rows))
        self.assertEqual(result[0], a.Target('192.0.2.1', 'Switch', 'Lab'))

    def test_excel_formula_rejected(self):
        with self.assertRaises(ValueError):
            a.parse_excel(self.workbook([['IP'], ['="192.0.2.1"']]))

    def test_excel_invalid_batch_rejected(self):
        with self.assertRaises(ValueError):
            a.parse_excel(self.workbook([['IP'], ['192.0.2.1'], ['bad']]))
        self.assertFalse(a.load_targets())

    def test_excel_duplicate_header(self):
        with self.assertRaises(ValueError):
            a.parse_excel(self.workbook([['IP', 'IP Address'], ['192.0.2.1', '192.0.2.2']]))

    def test_excel_xls_rejected(self):
        with self.assertRaises(ValueError):
            a.parse_excel(self.path / 'old.xls')

    def test_limits(self):
        with self.assertRaises(ValueError):
            a.Options(workers=100).validate()
        with self.assertRaises(ValueError):
            a.Options(ports='22; rm -rf').validate()
        with self.assertRaises(ValueError):
            a.Options(interval=1).validate()

    def test_no_authorization_no_run(self):
        self.engine = a.AutomationEngine(FakeAdapter)
        with self.assertRaises(PermissionError):
            self.engine.start([a.Target('192.0.2.1')], a.Options(), USER)
        self.assertFalse(self.engine.running)

    def test_viewer_denied(self):
        self.engine = a.AutomationEngine(FakeAdapter)
        with self.assertRaises(PermissionError):
            self.engine.start([a.Target('192.0.2.1')], a.Options(), {'role': 'Viewer'}, True)

    def test_missing_profile_before_inventory_write(self):
        self.engine = a.AutomationEngine(FakeAdapter)
        with self.assertRaises(ValueError):
            self.engine.start([a.Target('192.0.2.1', profile='Missing')], a.Options(), USER, True)
        with a.connect() as c:
            self.assertEqual(c.execute('SELECT COUNT(*) FROM network_devices').fetchone()[0], 0)

    def test_credential_encrypted_and_blank_keeps_secret(self):
        row = next(p for p in a.profiles() if p['name'] == 'Lab')
        self.assertNotIn('never-print-secret', row['community_enc'])
        a.save_profile('Lab', 'v2c', '', role='Admin')
        row2 = next(p for p in a.profiles() if p['name'] == 'Lab')
        self.assertEqual(row['community_enc'], row2['community_enc'])
        with self.assertRaises(PermissionError):
            a.save_profile('x', 'off', role='Operator')

    def test_shared_inventory_idempotent(self):
        t = a.Target('192.0.2.1', 'Important name')
        first = a.upsert_device(t)
        self.assertEqual(first, a.upsert_device(a.Target(t.ip, 'New name')))
        with a.connect() as c:
            self.assertEqual(c.execute('SELECT COUNT(*) FROM devices').fetchone()[0], 1)
            self.assertEqual(c.execute('SELECT name FROM network_devices').fetchone()[0], 'Important name')

    def test_ping_failure_does_not_skip_snmp(self):
        rows = self.run_job()
        statuses = {r['task']: r['status'] for r in rows}
        self.assertEqual(statuses['ping'], 'WARN')
        self.assertEqual(statuses['snmp'], 'OK')
        self.assertEqual(statuses['interfaces'], 'OK')
        self.assertEqual(statuses['topology'], 'OK')
        self.assertEqual(statuses['resources'], 'WARN')  # CPU yes, model-specific RAM unavailable.

    def test_snmp_failure_isolated_and_redacted(self):
        rows = self.run_job([a.Target('192.0.2.1'), a.Target('192.0.2.2')])
        self.assertTrue(any(r['ip'] == '192.0.2.2' and r['task'] == 'snmp' and r['status'] == 'ERROR' for r in rows))
        self.assertTrue(any(r['ip'] == '192.0.2.1' and r['task'] == 'snmp' and r['status'] == 'OK' for r in rows))
        self.assertNotIn('never-print-secret', json.dumps(rows))
        self.assertEqual(self.engine.snapshot()['done'], 2)

    def test_no_credential_skip_not_success(self):
        rows = self.run_job(default_profile='Default')
        self.assertTrue(any(r['task'] == 'snmp' and r['status'] == 'SKIP' for r in rows))
        self.assertTrue(any(r['task'] == 'tcp' and r['status'] == 'OK' for r in rows))

    def test_stale_v3_still_runs_ping_tcp(self):
        with a.connect() as c:
            c.execute("UPDATE autoip_profiles SET mode='v3',snmpv3_id=999999 WHERE name='Lab'")
        rows = self.run_job()
        self.assertTrue(any(r['task'] == 'tcp' and r['status'] == 'OK' for r in rows))
        self.assertTrue(any(r['task'] == 'snmp' and r['status'] == 'SKIP' for r in rows))

    def test_corrupt_community_still_runs_ping_tcp(self):
        with a.connect() as c:
            c.execute("UPDATE autoip_profiles SET community_enc='bad-token' WHERE name='Lab'")
        rows = self.run_job()
        self.assertTrue(any(r['task'] == 'tcp' and r['status'] == 'OK' for r in rows))
        self.assertTrue(any(r['task'] == 'snmp' and r['status'] == 'SKIP' for r in rows))

    def test_backup_off_by_default(self):
        rows = self.run_job()
        self.assertFalse(any(r['task'] == 'backup' for r in rows))

    def test_backup_missing_credentials_skip(self):
        rows = self.run_job(backup=True)
        self.assertTrue(any(r['task'] == 'backup' and r['status'] == 'SKIP' for r in rows))

    def test_concurrency_bounded(self):
        self.run_job([a.Target('192.0.2.' + str(i)) for i in range(10, 18)], workers=2)
        self.assertLessEqual(FakeAdapter.max_active, 2)
        self.assertEqual(FakeAdapter.active, 0)

    def test_double_start_and_stop(self):
        class Slow(FakeAdapter):
            def ping(self):
                self.stop.wait(0.5)
                self.check()
                return super().ping()
        self.engine = a.AutomationEngine(Slow)
        self.engine.start([a.Target('192.0.2.1')], a.Options(default_profile='Lab', report=False), USER, True)
        with self.assertRaises(RuntimeError):
            self.engine.start([a.Target('192.0.2.1')], a.Options(), USER, True)
        self.engine.stop()
        self.engine.thread.join(5)
        self.assertFalse(self.engine.running)
        self.assertEqual(self.engine.snapshot()['status'], 'Stopped')

    def test_repeat_wait_stop(self):
        self.engine = a.AutomationEngine(FakeAdapter)
        self.engine.start([a.Target('192.0.2.1')], a.Options(default_profile='Lab', report=False, alerts=False, repeat=True, interval=30), USER, True)
        deadline = time.monotonic() + 6
        while time.monotonic() < deadline and self.engine.snapshot()['status'] != 'Waiting':
            time.sleep(.02)
        self.assertEqual(self.engine.snapshot()['status'], 'Waiting')
        self.engine.stop()
        self.engine.thread.join(3)
        self.assertFalse(self.engine.running)
        self.assertEqual(len(a.run_history()), 1)

    def test_interrupted_run_recovery(self):
        with a.connect() as c:
            c.execute("INSERT INTO autoip_runs(id,status) VALUES('unfinished','Running')")
        self.engine = a.AutomationEngine(FakeAdapter)
        self.assertEqual(a.run_history()[0]['status'], 'Interrupted')
        self.assertFalse(self.engine.running)

    def test_counter_rates_and_reset(self):
        self.assertEqual(a.rate_from_counter(100, 1100, 10, 64, 1e9), 800)
        self.assertIsNone(a.rate_from_counter(1100, 100, 10, 64, 1e9))
        self.assertIsNone(a.rate_from_counter(100, 1100, 10, 64, 1e9, reset=True))
        self.assertIsNone(a.rate_from_counter(100, 1100, 300, 32, 1e9))
        self.assertIsNone(a.rate_from_counter(None, 1100, 10, 64, 1e9))

    def test_global_alert_postprocess(self):
        rows = self.run_job(alerts=True)
        self.assertTrue(any(r['task'] == 'alerts' and r['status'] == 'OK' for r in rows), rows)

    def test_report_and_formula_safety(self):
        rows = self.run_job(report=True)
        run_id = self.engine.snapshot()['run_id']
        self.engine._record(run_id, '192.0.2.1', 'snmp', 'OK', '=HYPERLINK("bad")')
        path = a.export_report(run_id, self.path / 'result.xlsx')
        wb = load_workbook(path, data_only=False)
        self.assertEqual(wb.sheetnames, ['Summary', 'Results'])
        self.assertFalse(any(cell.data_type == 'f' for ws in wb for row in ws for cell in row))
        self.assertEqual(wb['Results'].cell(wb['Results'].max_row, 4).value[0], "'")
        wb.close()
        self.assertTrue(Path(self.engine.snapshot()['report_path']).exists())
        self.assertNotIn('never-print-secret', path.read_bytes().decode('latin1'))

    def test_email_report_forces_export_and_uses_notification_smtp(self):
        import modules.advanced_pages as adv
        adv.set_setting('smtp_host', 'smtp.example.test')
        adv.set_setting('smtp_port', '587')
        adv.set_setting('smtp_user', 'sender@example.test')
        adv.set_setting('smtp_to', 'default@example.test')
        with patch.object(adv, '_load_secret', return_value='secret'), patch.object(adv, 'send_email') as mail:
            rows = self.run_job(email_report_after_run=True, report_email_to='receiver@gmail.com')
        snap = self.engine.snapshot()
        self.assertTrue(Path(snap['report_path']).exists())
        self.assertTrue(any(r['task'] == 'email_report' and r['status'] == 'OK' for r in rows), rows)
        self.assertEqual(mail.call_args.args[4], 'receiver@gmail.com')
        self.assertEqual(Path(mail.call_args.kwargs['attachment_path']), Path(snap['report_path']))

    def test_email_report_skips_when_smtp_not_configured(self):
        import modules.advanced_pages as adv
        adv.set_setting('smtp_host', '')
        adv.set_setting('smtp_to', '')
        rows = self.run_job(email_report_after_run=True)
        self.assertTrue(any(r['task'] == 'email_report' and r['status'] == 'SKIP' for r in rows), rows)

    def test_unknown_driver_prefix_boundary(self):
        did = a.upsert_device(a.Target('192.0.2.1'))
        self.assertIsNone(a.select_driver(did, '', '1.3.6.1.4.1.999.1'))

    def test_manual_driver_preserved(self):
        did = a.upsert_device(a.Target('192.0.2.1'))
        with a.connect() as c:
            driver = c.execute("SELECT id FROM vendor_drivers WHERE vendor='MikroTik'").fetchone()[0]
            c.execute('INSERT INTO device_driver_assignments(device_id,driver_id,source) VALUES(?,?,?)', (did, driver, 'manual'))
        self.assertEqual(a.select_driver(did, 'Cisco IOS', '1.3.6.1.4.1.9.1')['vendor'], 'MikroTik')

    def test_local_ping_error_not_offline(self):
        adapter = a.NetworkAdapter('127.0.0.1', {'mode': 'off'}, a.Options(), threading.Event())
        result = type('Result', (), {'returncode': 2, 'stdout': '', 'stderr': 'permission denied'})()
        with patch('subprocess.run', return_value=result):
            with self.assertRaises(RuntimeError):
                adapter.ping()

    def test_ssh_allowlist_blocks_writes(self):
        adapter = a.NetworkAdapter('192.0.2.1', {'mode': 'off', 'ssh': {'id': 1}}, a.Options(), threading.Event())
        for cmd in ('reload', 'write memory', 'show running-config; reload'):
            with self.assertRaises(a.SkipTask):
                adapter.backup({'backup_command': cmd})


if __name__ == '__main__':
    unittest.main(verbosity=2)
